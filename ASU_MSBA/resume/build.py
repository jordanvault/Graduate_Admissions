from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEADER = ("Sukwoo Chung".upper(), "swchung2469@gmail.com | +82-10-3613-3315 | linkedin.com/in/sukwoo-chung-b62ba1157 | Seoul, Korea")
EDU = [
 ("University of Wisconsin–Madison, Madison, WI", "09/07 – 08/13",
  "Bachelor of Business Administration, Finance, Investment & Banking | GPA 3.46/4.00", []),
 ("National University of Singapore (NUS) Business School, Singapore", "08/12 – 12/12",
  "Semester exchange program", []),
]
EXP = [
 ("Samsung SDS, Pangyo, Korea", "09/24 – Present", "Senior Data Analyst", [
  "Lead a six-person Freight Volume Aggregation Improvement Task Force to standardize volume classification across six regions and 10 organizations, establishing a company-wide authoritative dataset",
  "Redesigned four freight-aggregation rules, recovering 28,306 FEU of previously misclassified 2025 volume across Vessel, Air, Rail, and Truck businesses",
  "Diagnosed Cello Square customer journey friction using Google Analytics and internal data, identifying fragmented rate-search paths and quote delays linked to customer drop-off and recommending improvements",
  "Developed customized dashboards and retention tracking for a cosmetics-focused express initiative targeting Korea-to-U.S./Japan demand, enabling customer acquisition and repeat-use monitoring"]),
 ("Korean Air, Seoul, Korea / Manila, Philippines", "01/14 – 08/24", None, []),
 (None, "04/23 – 08/24", "Project Manager – Group Sales Optimization & RM Strategy", [
  "Led global implementation of Group Sales Optimizer, coordinating PROS, JAL/ANA, and Korean Air Sales & IT teams to automate group-booking availability and pricing decisions",
  "Built a reservation-misuse monitoring tool that identified 600+ cases annually and prevented an estimated $370K in annual revenue leakage, earning Korean Air's Quarterly IT Excellence Project Award"]),
 (None, "09/22 – 04/23", "Regional Sales Marketing Manager – Philippines Regional Office (Manila)", [
  "Led a seven-person sales team across Manila, Cebu, and Guam, driving reservation volume +340% YoY and restoring volume to 80% of the 2019 pre-COVID baseline",
  "Executed seven travel fairs with the Philippine Travel Agencies Association and Korea Tourism Organization, generating $2.1M in revenue"]),
 (None, "04/21 – 09/22", "Revenue Management & Pricing Systems Manager", [
  "Applied logistic regression and random forest models in Python to segment route-level COVID recovery patterns and inform demand forecasting and dynamic pricing decisions",
  "Evaluated Amadeus RMS as a competitive alternative during PROS contract renewal, contributing to $3.32M in five-year cost savings while mentoring four junior analysts"]),
 (None, "01/17 – 04/21", "Passenger Division Data Analyst & System Administrator", [
  "Integrated 100M+ ticketing and flown records with Delta Air Lines' passenger data in nine months, enabling the launch of the transpacific joint venture",
  "Renegotiated the IATA DDS data contract during COVID, securing $390K in cost savings; delivered 74 hours of Data/BI training in English and Korean"]),
 (None, "01/14 – 12/16", "Sales Support Assistant Manager & Sales Representative", [
  "Analyzed sales performance across 780+ Korean travel agencies using Oracle BI and supported the Mode Tour account across Southeast Asia and Americas markets"]),
]
AWARDS = [
 "Korean Air Quarterly IT Excellence Project Award — for the reservation-misuse monitoring tool (600+ cases identified, ~$370K annual revenue leakage prevented)",
 "PMP – Project Management Professional, Project Management Institute (2026)",
 "CPIM – Certified in Planning and Inventory Management, ASCM (2025)",
 "Mathematics for Machine Learning: Linear Algebra, Imperial College London (Coursera)",
]
ORGS = "[Professional organizations / volunteer / community involvement: 있으면 추가, 없으면 이 줄 삭제]"
VENTURE = ("Founder/Product Owner, Aimollae Book ([MM/YY] – Present): designed and commercialized a phone-hiding book for parents, "
           "selling 2,000+ units and generating approximately $25K in revenue; preparing U.S. market launch")
MIL = "Republic of Korea Army, mandatory military service — 09/09 – 07/11"
SKILLS = [
 ("Technical", "SQL, Python; Power BI, Oracle BI, Google Analytics, Figma; PROS O&D RMS/RTDP, Amadeus RMS, DDS, MIDT, CIRIUM; Claude Code for AI-assisted development"),
 ("Languages", "Korean (Native), Japanese (Fluent), Chinese (Basic)"),
 ("Global experience", "Worked with U.S., Japanese, and Philippine partners (Delta Air Lines, PROS, JAL/ANA); led sales and marketing for Korean Air's Philippines regional office (Manila, Cebu, Guam)"),
]
SALARY = [
 ("Samsung SDS (09/24 – Present)", "[연봉 입력: KRW ______ / 약 USD ______]"),
 ("Korean Air (01/14 – 08/24)", "[재직 종료 시점 연봉 입력, 또는 연도별 범위]"),
]

def hr(p):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); e=OxmlElement('w:bottom')
    for k,v in (('val','single'),('sz','6'),('space','1'),('color','000000')): e.set(qn('w:'+k),v)
    b.append(e); pPr.append(b)

d=Document()
s=d.sections[0]; s.left_margin=s.right_margin=Inches(0.7); s.top_margin=s.bottom_margin=Inches(0.6)
st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(9.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'),'Arial')
st.paragraph_format.space_after=Pt(0); st.paragraph_format.space_before=Pt(0)

def para(text='', bold=False, italic=False, size=None, align=None, after=0, before=0):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if align: p.alignment=align
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before); return p
def head(t):
    p=para(t.upper(), bold=True, size=10.5, before=7, after=2); hr(p)
def line(left, right, bold=True, italic=False):
    p=d.add_paragraph(); p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), alignment=2)
    r=p.add_run(left); r.bold=bold; r.italic=italic
    p.add_run('\t'+right); return p
def bullet(t):
    p=d.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(0.22)
    p.paragraph_format.space_after=Pt(1); p.add_run(t); return p

para(HEADER[0], bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
para(HEADER[1], size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
head('Education')
for sch,dt,deg,_ in EDU:
    line(sch,dt); para(deg, italic=True, after=2)
head('Professional Experience')
for org,dt,title,bl in EXP:
    if org and title is None:
        line(org,dt); continue
    if org:
        line(org,dt); para(title, italic=True)
    else:
        line(title,dt,bold=True,italic=False)
    for b in bl: bullet(b)
head('Honors, Awards & Certifications')
for a in AWARDS: bullet(a)
head('Entrepreneurship, Service & Community')
bullet(VENTURE); bullet(MIL); bullet(ORGS)
head('Skills & Global Experience')
for k,v in SKILLS:
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(1); r=p.add_run(k+': '); r.bold=True; p.add_run(v)
head('Salary History')
for k,v in SALARY:
    p=d.add_paragraph(); r=p.add_run(k+': '); r.bold=True; p.add_run(v)
d.save('Sukwoo_Chung_Resume_ASU_MSBA.docx')

# markdown twin
md=["# SUKWOO CHUNG", HEADER[1], "", "## EDUCATION"]
for sch,dt,deg,_ in EDU: md+= [f"**{sch}** — {dt}", f"*{deg}*", ""]
md.append("## PROFESSIONAL EXPERIENCE")
for org,dt,title,bl in EXP:
    if org and title is None: md+=[f"\n**{org}** — {dt}"]; continue
    md.append(f"\n**{org}** — {dt}\n*{title}*" if org else f"\n*{title}* — {dt}")
    md+= [f"- {b}" for b in bl]
md+=["","## HONORS, AWARDS & CERTIFICATIONS"]+[f"- {a}" for a in AWARDS]
md+=["","## ENTREPRENEURSHIP, SERVICE & COMMUNITY",f"- {VENTURE}",f"- {MIL}",f"- {ORGS}","","## SKILLS & GLOBAL EXPERIENCE"]
md+=[f"- **{k}:** {v}" for k,v in SKILLS]
md+=["","## SALARY HISTORY"]+[f"- **{k}:** {v}" for k,v in SALARY]
open('resume.md','w').write("\n".join(md)+"\n")
